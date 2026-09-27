from __future__ import annotations

from src.services.parsing_adapters.models import ParsedDocument, ParsedPage


def _table_to_text(table) -> str:
    rows: list[str] = []
    for row in table.rows:
        seen: set[int] = set()
        cells: list[str] = []
        for cell in row.cells:
            key = id(cell._tc)
            if key in seen:
                continue
            seen.add(key)
            if t := cell.text.strip():
                cells.append(t)
            for nested in cell.tables:
                if nt := _table_to_text(nested):
                    cells.append(nt)
        if cells:
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _iter_block_items(parent):
    from docx.document import Document as DocxDocument
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = parent.element.body if isinstance(parent, DocxDocument) else parent._element
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def _story_parts(container) -> list[str]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parts: list[str] = []
    for block in _iter_block_items(container):
        if isinstance(block, Paragraph):
            if t := (block.text or "").strip():
                parts.append(t)
        elif isinstance(block, Table):
            if t := _table_to_text(block):
                parts.append(t)
    return parts


def _header_footer_parts(doc) -> list[str]:
    # 单节异常跳过，避免整份 docx 失败
    seen: set[str] = set()
    out: list[str] = []
    for section in doc.sections:
        for part in (section.header, section.footer):
            try:
                block = "\n".join(_story_parts(part)).strip()
            except Exception:
                continue
            if block and block not in seen:
                seen.add(block)
                out.append(block)
    return out


def _textbox_parts(doc) -> list[str]:
    from docx.oxml.ns import qn

    parts: list[str] = []
    for txbx in doc.element.body.iter(qn("w:txbxContent")):
        texts = [(n.text or "").strip() for n in txbx.iter(qn("w:t")) if (n.text or "").strip()]
        if texts:
            parts.append("\n".join(texts))
    return parts


def _parse_docx(path: str) -> ParsedDocument:
    from docx import Document

    doc = Document(path)
    parts = _header_footer_parts(doc) + _story_parts(doc) + _textbox_parts(doc)
    text = "\n\n".join(parts)
    return ParsedDocument([ParsedPage(page=None, text=text)] if text else [])
