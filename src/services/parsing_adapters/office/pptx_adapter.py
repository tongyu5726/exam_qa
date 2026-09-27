from __future__ import annotations

import logging
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from src.services.parsing_adapters.commands import _run_capture
from src.services.parsing_adapters.models import ParsedDocument, ParsedPage
from src.services.parsing_adapters.office.soffice import _find_soffice
from src.services.parsing_adapters.pdf.pymupdf_adapter import _parse_pdf_pymupdf4llm
from src.services.parsing_adapters.pdf.strategy import _parse_pdf

logger = logging.getLogger(__name__)


def _shape_texts(shape) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
        return [t for child in shape.shapes for t in _shape_texts(child)]

    texts: list[str] = []
    if getattr(shape, "has_table", False):
        rows = []
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            texts.append("\n".join(rows))
    if getattr(shape, "has_text_frame", False):
        if t := shape.text_frame.text.strip():
            texts.append(t)
    elif hasattr(shape, "text") and (t := (shape.text or "").strip()):
        texts.append(t)
    return texts


def _convert_pptx_to_pdf_soffice(path: str) -> Path | None:
    """LibreOffice headless 将 PPTX 转为 PDF，可覆盖 SmartArt/图表/母版文本。"""
    soffice = _find_soffice()
    if not soffice:
        return None

    out_dir = Path(tempfile.mkdtemp(prefix="exam_pptx_"))
    profile_dir = Path(tempfile.mkdtemp(prefix="exam_lo_profile_"))
    dest: Path | None = None
    try:
        user_install = "file:///" + str(profile_dir).replace("\\", "/")
        proc, _stdout, stderr = _run_capture(
            [
                soffice,
                "--headless",
                "--norestore",
                "--nolockcheck",
                "--convert-to",
                "pdf",
                "--outdir",
                str(out_dir),
                f"-env:UserInstallation={user_install}",
                path,
            ],
            timeout=180,
        )
        if proc.returncode != 0:
            logger.warning(
                "LibreOffice 退出码 %s: %s", proc.returncode, stderr.strip()[:300]
            )
            return None
        converted = out_dir / f"{Path(path).stem}.pdf"
        if not converted.is_file():
            return None
        fd, dest_name = tempfile.mkstemp(suffix=".pdf", prefix="exam_pptx_")
        os.close(fd)
        dest = Path(dest_name)
        shutil.copy2(converted, dest)
        return dest
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError) as e:
        logger.warning("LibreOffice PPTX 转 PDF 失败: %s", e)
        if dest is not None:
            dest.unlink(missing_ok=True)
        return None
    finally:
        shutil.rmtree(out_dir, ignore_errors=True)
        shutil.rmtree(profile_dir, ignore_errors=True)


def _convert_pptx_to_pdf_powerpoint(path: str) -> Path | None:
    """Windows：无 LibreOffice 时用本机 PowerPoint COM 导出 PDF（ppSaveAsPDF=32）。"""
    if os.name != "nt":
        return None
    src = str(Path(path).resolve())
    fd, dest_name = tempfile.mkstemp(suffix=".pdf", prefix="exam_pptx_")
    os.close(fd)
    dest = Path(dest_name)
    dest.unlink(missing_ok=True)
    src_ps = src.replace("'", "''")
    dest_ps = str(dest.resolve()).replace("'", "''")
    ps = (
        "$ErrorActionPreference='Stop'; "
        "$pp=New-Object -ComObject PowerPoint.Application; "
        "try { "
        f"$pres=$pp.Presentations.Open('{src_ps}',-1,0,0); "
        f"$pres.SaveAs([ref]'{dest_ps}',[ref]32); "
        "$pres.Close(); "
        "} finally { "
        "$pp.Quit(); "
        "[System.Runtime.InteropServices.Marshal]::ReleaseComObject($pp)|Out-Null "
        "}"
    )
    try:
        subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", ps],
            check=True,
            capture_output=True,
            text=True,
            timeout=180,
        )
        if dest.is_file() and dest.stat().st_size > 0:
            return dest
        dest.unlink(missing_ok=True)
        return None
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError) as e:
        logger.warning("PowerPoint COM 导出 PDF 失败: %s", e)
        dest.unlink(missing_ok=True)
        return None


def _convert_pptx_to_pdf(path: str) -> Path | None:
    return _convert_pptx_to_pdf_soffice(path) or _convert_pptx_to_pdf_powerpoint(path)


def _parse_pptx_direct(path: str) -> ParsedDocument:
    from pptx import Presentation

    pages: list[ParsedPage] = []
    for i, slide in enumerate(Presentation(path).slides, 1):
        texts = [t for shape in slide.shapes for t in _shape_texts(shape)]
        if slide.has_notes_slide:
            if notes := slide.notes_slide.notes_text_frame.text.strip():
                texts.append(f"[备注]\n{notes}")
        if texts:
            pages.append(ParsedPage(page=i, text="\n".join(texts)))
    return ParsedDocument(pages)


def _tesseract_available() -> bool:
    if shutil.which("tesseract"):
        return True
    if os.name == "nt" and Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe").is_file():
        return True
    return False


def _clean_pdf_markdown(text: str) -> str:
    """清理 pymupdf4llm 输出的 markdown 杂质（图片占位符、标题标记等）。"""
    text = re.sub(r"\*\*==>.*?<==\*\*", "", text, flags=re.S)
    text = re.sub(r"^\s*#+\s*$", "", text, flags=re.M)
    text = re.sub(r"^\s*#{1,6}\s+", "", text, flags=re.M)
    text = text.replace("**", "").replace("__", "")
    text = re.sub(r"^\s*\|?[\s:|-]+\|?\s*$", "", text, flags=re.M)
    lines = [ln.rstrip() for ln in text.splitlines() if ln.strip()]
    return "\n".join(lines).strip()


def _merge_pages_by_length(a: list[ParsedPage], b: list[ParsedPage]) -> ParsedDocument:
    """同一 PPT 的两种提取结果按页择优合并（PDF 页号与 PPT 页号一致）。"""
    pages: list[ParsedPage] = []
    for i in range(max(len(a), len(b))):
        pa = a[i] if i < len(a) else None
        pb = b[i] if i < len(b) else None
        if pa and pb:
            pages.append(pa if len(pa.text) >= len(pb.text) else pb)
        else:
            pages.append(pa or pb)
    return ParsedDocument([p for p in pages if p.text and p.text.strip()])


def _parse_pptx(path: str) -> ParsedDocument:
    """优先转 PDF 走 pymupdf4llm 管线，清理杂质后与 python-pptx 按页择优合并；
    图片型 PPT 自动尝试 OCR（需系统安装 Tesseract）。"""
    direct = _parse_pptx_direct(path)
    pdf_tmp = _convert_pptx_to_pdf(path)
    if not pdf_tmp:
        return direct
    try:
        doc = _parse_pdf(str(pdf_tmp))
        pdf_pages = [
            ParsedPage(page=p.page, text=_clean_pdf_markdown(p.text))
            for p in doc.pages
            if _clean_pdf_markdown(p.text).strip()
        ]
        direct_chars = sum(len(p.text) for p in direct.pages)
        pdf_chars = sum(len(p.text) for p in pdf_pages)
        if pdf_chars < max(direct_chars, 1) * 0.5 and _tesseract_available():
            try:
                ocr_doc = _parse_pdf_pymupdf4llm(str(pdf_tmp), force_ocr=True)
                ocr_pages = [
                    ParsedPage(page=p.page, text=_clean_pdf_markdown(p.text))
                    for p in ocr_doc.pages
                    if _clean_pdf_markdown(p.text).strip()
                ]
                ocr_chars = sum(len(p.text) for p in ocr_pages)
                if ocr_chars > pdf_chars:
                    logger.info("PPTX 图片 OCR 生效: %s -> %d 字", Path(path).name, ocr_chars)
                    pdf_pages = ocr_pages
            except Exception as e:
                logger.warning("PPTX 图片 OCR 失败: %s", e)
        return _merge_pages_by_length(direct.pages, pdf_pages)
    except Exception as e:
        logger.warning("PPTX 转 PDF 解析失败，回退 python-pptx: %s", e)
        return direct
    finally:
        pdf_tmp.unlink(missing_ok=True)
