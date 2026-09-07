"""文档解析：PDF / Office / 纯文本 → ParsedDocument。"""

from __future__ import annotations

import base64
import hashlib
import json
import logging
import os
import re
import shlex
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from src.config import config
from src.exceptions import BadRequestException, UnsupportedFormatException

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".doc", ".docx", ".pptx"}


def _run_capture(args: list[str], *, timeout: float | None = None, encoding: str = "utf-8"):
    """subprocess.run 的 Windows 安全版：stdout/stderr 写临时文件而非管道。

    子进程若再派生子进程（COM dllhost / MinerU 子进程）会继承管道句柄，
    导致 communicate() 在超时后仍永久阻塞；临时文件无句柄继承问题。
    """
    import tempfile as _tf

    with _tf.TemporaryFile() as _out, _tf.TemporaryFile() as _err:
        proc = subprocess.run(args, stdout=_out, stderr=_err, timeout=timeout)
        _out.seek(0)
        _err.seek(0)
        stdout = _out.read().decode(encoding, errors="replace")
        stderr = _err.read().decode(encoding, errors="replace")
    return proc, stdout, stderr


def _command_error_tail(stdout: str, stderr: str, limit: int = 1200) -> str:
    """外部 CLI 的最后几行通常才包含 traceback / root cause，避免日志只留下启动信息。"""
    output = (stderr or stdout).strip()
    return output[-limit:] if output else "<无错误输出>"


@dataclass
class ParsedPage:
    page: int | None  # PDF/PPT 1-based；纯文本/docx 为 None
    text: str
    # PDF 文档自身定义的页签；和物理页 page、教材中写的 P26 等引用严格分开。
    pdf_page_label: str = ""


@dataclass
class ParsedBlock:
    """MinerU 结构化 block：结构还原的最小单元。"""

    block_type: str  # text|title|table|formula|formula_inline|image|image_caption|table_caption|figure_caption|header|footer|...
    text: str  # 可检索文本（表格/公式/图片摘要已转文本）
    page: int | None  # 1-based
    level: int = 0  # 标题层级 1-6；非标题为 0
    html: str = ""  # 表格原始 HTML
    latex: str = ""  # 公式 LaTeX
    image_path: str = ""  # 图片原始文件路径（多模态摘要用）
    caption: str = ""  # 图片/表格说明
    order: int = 0  # 原始顺序（跨页排序用）
    bbox: tuple | None = None
    pdf_page_label: str = ""
    textbook_references: tuple[str, ...] = ()
    # content | annotation | reference | noise。噪声不进入向量库，批注保留为独立证据。
    content_role: str = "content"


@dataclass
class ParsedDocument:
    pages: list[ParsedPage]
    blocks: list[ParsedBlock] = field(default_factory=list)
    # 解析链路的轻量可追溯信息：不增加模型调用，随切片写入向量 metadata。
    parser_name: str = ""
    parse_quality: float = 0.0

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text.strip())

    @property
    def has_blocks(self) -> bool:
        return bool(self.blocks)


@dataclass(frozen=True)
class PDFParseQuality:
    """用于选择 PDF 解析候选结果的轻量质量信号（不依赖模型调用）。"""

    score: float
    page_coverage: float
    chars_per_page: float
    structure_ratio: float
    expected_pages: int
    parsed_pages: int


def _pdf_page_count(path: str) -> int:
    """读取物理页数；失败时返回 0，质量评估仍可继续。"""
    try:
        import fitz

        doc = fitz.open(path)
        try:
            return doc.page_count
        finally:
            doc.close()
    except Exception:
        return 0


def _assess_pdf_quality(
    doc: ParsedDocument | None, *, expected_pages: int = 0
) -> PDFParseQuality:
    """以文本密度、页覆盖与结构还原评估候选结果，范围固定为 0～1。"""
    if doc is None:
        return PDFParseQuality(0.0, 0.0, 0.0, 0.0, expected_pages, 0)

    pages_with_text = [page for page in doc.pages if page.text.strip()]
    chars = len(re.sub(r"\s", "", doc.full_text))
    parsed_pages = len(pages_with_text)
    # Markdown 回退只有一页 None，不能把它误当作仅覆盖原 PDF 的第一页。
    if expected_pages and any(page.page is None for page in pages_with_text) and chars >= 80:
        page_coverage = 1.0
    else:
        denominator = expected_pages or max(len(doc.pages), 1)
        page_coverage = min(parsed_pages / denominator, 1.0)
    chars_per_page = chars / max(expected_pages or parsed_pages, 1)
    text_density = min(chars_per_page / 160, 1.0)
    structure_ratio = 1.0 if doc.blocks else (0.65 if chars else 0.0)
    score = round(0.45 * text_density + 0.40 * page_coverage + 0.15 * structure_ratio, 4)
    return PDFParseQuality(
        score=score,
        page_coverage=round(page_coverage, 4),
        chars_per_page=round(chars_per_page, 2),
        structure_ratio=structure_ratio,
        expected_pages=expected_pages,
        parsed_pages=parsed_pages,
    )


def _log_pdf_quality(source: str, quality: PDFParseQuality) -> None:
    logger.info(
        "PDF 候选质量: %s score=%.2f coverage=%.0f%% chars/page=%.0f structure=%.0f%%",
        source,
        quality.score,
        quality.page_coverage * 100,
        quality.chars_per_page,
        quality.structure_ratio * 100,
    )


def _page_num(meta: dict) -> int | None:
    raw = meta.get("page_number") or meta.get("page")
    if raw is None:
        return None
    try:
        return int(raw)
    except (TypeError, ValueError):
        return None


def _pages_from_pymupdf4llm(raw) -> ParsedDocument | None:
    if isinstance(raw, str):
        text = raw.strip()
        return ParsedDocument([ParsedPage(page=None, text=text)]) if text else None
    pages = []
    for chunk in raw:
        text = (chunk.get("text") or "").strip()
        if text:
            pages.append(ParsedPage(page=_page_num(chunk.get("metadata") or {}), text=text))
    return ParsedDocument(pages) if pages else None


def _pymupdf4llm_kwargs(*, force_ocr: bool | None = None) -> dict:
    p = config.parsing
    return {
        "page_chunks": True,
        "use_ocr": p.pdf_use_ocr,
        "force_ocr": p.pdf_force_ocr if force_ocr is None else force_ocr,
        "ocr_language": p.pdf_ocr_language,
    }


def _parse_pdf_pymupdf4llm(path: str, *, force_ocr: bool | None = None) -> ParsedDocument | None:
    import pymupdf4llm

    # pymupdf4llm 调用 Tesseract 时从进程环境读取 TESSDATA_PREFIX。
    # 配置为空不覆盖用户/系统已有设置。
    tessdata_prefix = config.parsing.tessdata_prefix
    if tessdata_prefix:
        os.environ["TESSDATA_PREFIX"] = tessdata_prefix
    return _pages_from_pymupdf4llm(
        pymupdf4llm.to_markdown(path, **_pymupdf4llm_kwargs(force_ocr=force_ocr))
    )


def _parse_pdf_fitz(path: str) -> ParsedDocument:
    import fitz

    doc = fitz.open(path)
    try:
        pages = []
        for i, page in enumerate(doc, 1):
            text = page.get_text().strip()
            if not text:
                continue
            try:
                label = page.get_label() or ""
            except Exception:
                label = ""
            pages.append(ParsedPage(page=i, text=text, pdf_page_label=label))
        return ParsedDocument(pages)
    finally:
        doc.close()


def _mineru_available(cmd: str) -> bool:
    """MinerU CLI 是否可用（PATH 可找到）。"""
    return bool(shutil.which(cmd))


def _pdf_kind(path: str) -> str:
    """抽样判断 PDF 类型：text（原生文本）| scanned（扫描件/图片型）| mixed。

    原生文本 PDF 直接提取文本；扫描件/混合型交给 OCR / MinerU。
    采样首页、次页、中间页、末页，避免只按首字符数误判。
    """
    try:
        import fitz

        doc = fitz.open(path)
        try:
            n = doc.page_count
            if n == 0:
                return "scanned"
            total = 0
            sampled = 0
            for i in sorted({0, 1, n // 2, n - 1}):
                if i >= n:
                    continue
                total += len(re.sub(r"\s", "", doc[i].get_text() or ""))
                sampled += 1
            if sampled == 0 or total < 30:
                return "scanned"
            avg = total / sampled
            if avg < 30:
                return "scanned"
            return "mixed" if avg < 200 else "text"
        finally:
            doc.close()
    except Exception:
        return "scanned"


def _html_table_to_text(html: str) -> str:
    """把 MinerU 的表格 HTML 转成可检索的纯文本（单元格用 | 分隔）。"""
    if not html:
        return ""
    text = re.sub(r"<t[dh][^>]*>", " | ", html, flags=re.I)
    text = re.sub(r"</t[r]>", "\n", text, flags=re.I)
    text = re.sub(r"</t[dh]>", "", text, flags=re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"&nbsp;?", " ", text, flags=re.I)
    lines = [re.sub(r"\s*[|]\s*$", "", ln).strip() for ln in text.splitlines()]
    return "\n".join(ln for ln in lines if ln.strip())


def _table_headers(html: str) -> str:
    """从表格 HTML 提取列头（<th>），写入结构化 metadata。"""
    if not html:
        return ""
    heads = re.findall(r"<t[h][^>]*>(.*?)</t[h]>", html, flags=re.I | re.S)
    cleaned = [" ".join(re.sub(r"<[^>]+>", " ", h).split()) for h in heads]
    return " | ".join(c for c in cleaned if c)


def _normalize_formula_latex(value: object) -> str:
    """将不同解析器的公式字段归一为裸 LaTex，避免重复包裹分隔符。"""
    latex = str(value or "").strip()
    if latex.startswith("$$") and latex.endswith("$$") and len(latex) >= 4:
        return latex[2:-2].strip()
    if latex.startswith("\\[") and latex.endswith("\\]") and len(latex) >= 4:
        return latex[2:-2].strip()
    if latex.startswith("\\(") and latex.endswith("\\)") and len(latex) >= 4:
        return latex[2:-2].strip()
    if latex.startswith("$") and latex.endswith("$") and len(latex) >= 2:
        return latex[1:-1].strip()
    return latex


def _formula_latex(block: dict) -> str:
    """兼容 MinerU / 其他解析器的 latex、math_content、text 等公式字段。"""
    for key in ("latex", "rec_formula", "formula", "formula_text", "math_content", "math", "content", "text"):
        value = _normalize_formula_latex(block.get(key))
        if value:
            return value
    return ""


def _paddle_result_payloads(result: object) -> list[dict]:
    """兼容 PaddleOCR 3.x Result 的 json/dict 形态，提取可能含公式的字典。"""
    payload = getattr(result, "json", result)
    if callable(payload):
        payload = payload()
    if isinstance(payload, str):
        try:
            payload = json.loads(payload)
        except json.JSONDecodeError:
            return []
    found: list[dict] = []

    def walk(value: object, inherited_page: int | None = None) -> None:
        if isinstance(value, list):
            for item in value:
                walk(item, inherited_page)
            return
        if not isinstance(value, dict):
            return
        # PaddleOCR 的 page_index/page_idx 均为零基页码；项目内部与 PDF
        # 其余解析器统一使用一基页码。`page` 则视为已经归一化的一基值。
        if value.get("page_index") is not None:
            page = value.get("page_index")
            zero_based = True
        elif value.get("page_idx") is not None:
            page = value.get("page_idx")
            zero_based = True
        else:
            page = value.get("page")
            zero_based = False
        try:
            page = int(page) + 1 if zero_based else int(page)
        except (TypeError, ValueError):
            page = inherited_page
        latex = _formula_latex(value)
        # 只接受明确的公式输出键，防止把一般文字字段误写成 LaTeX。
        has_formula_key = any(key in value for key in ("latex", "rec_formula", "formula", "formula_text"))
        if latex and has_formula_key:
            item = dict(value)
            item["_formula_latex"] = latex
            item["_page"] = page
            found.append(item)
        for child in value.values():
            if isinstance(child, (dict, list)):
                walk(child, page)

    walk(payload)
    return found


def _formula_bbox(raw: object) -> tuple | None:
    """兼容 PaddleOCR 的 bbox/coordinate/polygon 输出，统一成 x0,y0,x1,y1。"""
    if isinstance(raw, (list, tuple)) and len(raw) == 4 and all(
        isinstance(value, (int, float)) for value in raw
    ):
        return _as_bbox(raw)
    if isinstance(raw, (list, tuple)) and len(raw) >= 4:
        points = [point for point in raw if isinstance(point, (list, tuple)) and len(point) >= 2]
        if points:
            try:
                xs = [float(point[0]) for point in points]
                ys = [float(point[1]) for point in points]
                return min(xs), min(ys), max(xs), max(ys)
            except (TypeError, ValueError):
                return None
    return None


def _run_formula_worker(path: str) -> list[dict]:
    """在独立进程运行 PaddleOCR，隔离 Windows 下 Torch/Paddle 的 cuDNN DLL。"""
    p = config.parsing
    with tempfile.TemporaryDirectory(prefix="formula_worker_") as tmp:
        output = Path(tmp) / "payloads.json"
        args = [
            sys.executable,
            "-m",
            "src.services.formula_worker",
            "--input",
            str(Path(path).resolve()),
            "--output",
            str(output),
            "--device",
            p.formula_recognition_device,
            "--model",
            p.formula_recognition_model,
        ]
        if p.formula_recognition_enable_mkldnn:
            args.append("--enable-mkldnn")
        timeout = p.formula_recognition_timeout or None
        proc, stdout, stderr = _run_capture(args, timeout=timeout)
        if proc.returncode != 0:
            raise RuntimeError(
                "PaddleOCR 公式子进程退出码 "
                f"{proc.returncode}: {_command_error_tail(stdout, stderr)}"
            )
        if not output.is_file():
            raise RuntimeError(
                "PaddleOCR 公式子进程未产出 JSON: "
                + _command_error_tail(stdout, stderr)
            )
        payloads = json.loads(output.read_text(encoding="utf-8"))
        if not isinstance(payloads, list):
            raise RuntimeError("PaddleOCR 公式子进程输出格式错误")
        return [item for item in payloads if isinstance(item, dict)]


def _run_formula_pipeline(path: str) -> list[dict]:
    """运行公式识别；Windows GPU 使用子进程以避免 cuDNN DLL 版本冲突。"""
    p = config.parsing
    device = p.formula_recognition_device.strip().lower()
    if os.name == "nt" and (device == "auto" or device.startswith(("gpu", "cuda"))):
        logger.info("公式识别使用独立 GPU 子进程: device=%s", p.formula_recognition_device)
        return _run_formula_worker(path)

    # CPU Torch 与 Paddle 在当前 Windows 环境可按该顺序共存；GPU 路径必须走上面的隔离分支。
    if os.name == "nt":
        import torch  # noqa: F401

    from paddleocr import FormulaRecognitionPipeline
    from src.services.inference_device import resolve_paddle_device

    pipeline = FormulaRecognitionPipeline(
        device=resolve_paddle_device(p.formula_recognition_device),
        formula_recognition_model_name=p.formula_recognition_model,
        enable_mkldnn=p.formula_recognition_enable_mkldnn,
    )
    payloads: list[dict] = []
    for result in pipeline.predict(path):
        payloads.extend(_paddle_result_payloads(result))
    return payloads


def _enrich_document_with_formulas(doc: ParsedDocument | None, path: str) -> ParsedDocument | None:
    """对全文补偿公式块；失败时保留原解析结果，不中断文档入库。"""
    p = config.parsing
    if not doc or not doc.pages or not p.formula_recognition_enabled:
        return doc
    try:
        payloads = _run_formula_pipeline(path)
    except Exception as exc:
        logger.warning("全文公式识别失败，保留原解析文本: %s", exc)
        return doc

    if not payloads:
        logger.info("全文公式识别未发现可用 LaTeX: %s", Path(path).name)
        return doc

    # 回退解析页先显式建为文本块；否则 has_blocks=True 后结构化切片会只保留公式而丢正文。
    blocks = list(doc.blocks) or [
        _annotate_block(
            ParsedBlock(
                block_type="text",
                text=page.text,
                page=page.page,
                order=(page.page or index + 1) * 100000,
                pdf_page_label=page.pdf_page_label,
            )
        )
        for index, page in enumerate(doc.pages)
        if page.text.strip()
    ]
    per_page_order: dict[int | None, int] = {}
    formulas: list[ParsedBlock] = []
    existing_formula_ids = {
        re.sub(r"\s+", "", block.latex or _formula_latex(block.text))
        for block in blocks
        if block.block_type in {"formula", "formula_inline"}
    }
    for payload in payloads:
        page = payload.get("_page")
        try:
            page = int(page) if page is not None else None
        except (TypeError, ValueError):
            page = None
        per_page_order[page] = per_page_order.get(page, 0) + 1
        latex = str(payload["_formula_latex"])
        formula_id = re.sub(r"\s+", "", latex)
        if not formula_id or formula_id in existing_formula_ids:
            continue
        existing_formula_ids.add(formula_id)
        formula = _annotate_block(
            ParsedBlock(
                block_type="formula",
                text="$$" + latex + "$$",
                latex=latex,
                page=page,
                order=(page or 0) * 100000 + per_page_order[page],
                bbox=_formula_bbox(
                    payload.get("bbox") or payload.get("coordinate") or payload.get("polygon")
                ),
            )
        )
        formulas.append(formula)
    blocks.extend(formulas)
    blocks.sort(key=lambda block: (block.page or 0, block.order))
    for formula in formulas:
        for page in doc.pages:
            if page.page == formula.page:
                page.text = (page.text + "\n\n" + formula.text).strip()
                break
    doc.blocks = blocks
    logger.info("全文公式识别完成: %s, 新增 %d 条 LaTeX 公式", Path(path).name, len(formulas))
    return doc


_TEXTBOOK_REFERENCE_PATTERNS = (
    re.compile(r"(?<![A-Za-z0-9])(?:P|p)\s*\d{1,4}(?:\s*(?:例|Ex|T|习题)\s*\d+(?:\s*\([^)]{1,20}\))*)?"),
    re.compile(r"第\s*\d{1,4}\s*页"),
    re.compile(r"(?<![\u4e00-\u9fffA-Za-z0-9])(?:例|Ex|T)\s*\d+(?:\s*\([^)]{1,20}\))*", re.I),
)
_ANNOTATION_MARKER = re.compile(r"^(?:考点|注意|易错|方法|提示|结论|证明|思路)\s*[:：]", re.I)
_UI_NOISE = re.compile(
    r"^(?:\d{1,2}:\d{2}|[2-5]G|[0-9]{1,3}%|仅供参考|如有错误欢迎指正|"
    r"https?://\S+|[\w.+-]+@[\w.-]+\.[A-Za-z]{2,})$",
    re.I,
)


def _extract_textbook_references(text: str) -> tuple[str, ...]:
    """提取“教材内页码/题号”候选，不把它误作 PDF 物理页。"""
    found: list[str] = []
    for pattern in _TEXTBOOK_REFERENCE_PATTERNS:
        for match in pattern.finditer(text or ""):
            value = re.sub(r"\s+", " ", match.group(0)).strip()
            if value and value not in found:
                found.append(value)
    # P28 Ex2 同时会命中 P28 Ex2 与 Ex2；保留信息更完整的引用即可。
    return tuple(
        value
        for value in found
        if not any(value != candidate and value in candidate for candidate in found)
    )


def _content_role(block_type: str, text: str) -> str:
    """保守地标记版面角色；只滤明显噪声，红蓝批注等仍作为可检索证据保留。"""
    if block_type in {"header", "footer", "page_number", "footnote"}:
        return "noise"
    clean = " ".join((text or "").split())
    if clean and len(clean) <= 100 and _UI_NOISE.fullmatch(clean):
        return "noise"
    if _ANNOTATION_MARKER.match(clean):
        return "annotation"
    if clean and len(clean) <= 80 and _extract_textbook_references(clean):
        return "reference"
    return "content"


def _annotate_block(block: ParsedBlock) -> ParsedBlock:
    block.textbook_references = _extract_textbook_references(block.text)
    block.content_role = _content_role(block.block_type, block.text)
    return block


def _block_text(block: dict) -> str:
    """content_list 单个 block 的可检索文本（text/table/formula 按类型拼接）。"""
    btype = block.get("type") or ""
    if btype in ("formula", "formula_inline", "equation", "equation_interline", "inline_equation"):
        latex = _formula_latex(block)
        return f"$${latex}$$" if latex else ""
    if btype == "table":
        text = block.get("text") or ""
        html = block.get("html") or ""
        return text if text else _html_table_to_text(html)
    text = block.get("text") or ""
    if text:
        return text
    parts = []
    for ln in block.get("lines") or []:
        if isinstance(ln, dict):
            t = ln.get("text") or ln.get("content") or ""
        else:
            t = str(ln)
        if t.strip():
            parts.append(t.strip())
    return "\n".join(parts)


_CAPTION_TYPES = frozenset({"image_caption", "table_caption", "figure_caption"})


def _caption_text(block: dict) -> str:
    """图片/表格说明文本（兼容 text 与 lines 两种形态）。"""
    text = (block.get("text") or "").strip()
    if text:
        return text
    parts = []
    for ln in block.get("lines") or []:
        if isinstance(ln, dict):
            t = ln.get("text") or ln.get("content") or ""
        else:
            t = str(ln)
        if t.strip():
            parts.append(t.strip())
    return " ".join(parts)


def _as_bbox(raw) -> tuple | None:
    if not isinstance(raw, (list, tuple)) or len(raw) != 4:
        return None
    try:
        return tuple(float(v) for v in raw)
    except (TypeError, ValueError):
        return None


def _resolve_image_path(img_path: str | None, base_dir: Path | None) -> str:
    """MinerU 图片路径可能是相对 json 输出目录的，解析为绝对路径。"""
    if not img_path:
        return ""
    p = Path(img_path)
    if not p.is_absolute() and base_dir is not None:
        p = base_dir / p
    return str(p.resolve()) if p.exists() else ""


def _persist_parser_assets(
    document: ParsedDocument, *, work_dir: Path, source_path: str
) -> None:
    """在清理外部解析器临时目录前，持久化其中的图片资产。

    MinerU / MarkPDFdown 的 JSON 可能引用临时目录中的裁剪图。若不复制，入库阶段
    的视觉摘要必然读不到图片。仅复制位于本次 work_dir 下、且确实被 block 引用的文件。
    """
    asset_blocks = [b for b in document.blocks if b.image_path]
    if not asset_blocks:
        return
    try:
        source = Path(source_path)
        try:
            stamp = f"{source.resolve()}:{source.stat().st_size}:{source.stat().st_mtime_ns}"
        except OSError:
            stamp = str(source.resolve())
        target_dir = Path(config.storage.parsed_assets_dir).resolve() / hashlib.sha256(
            stamp.encode("utf-8", errors="replace")
        ).hexdigest()[:16]
        target_dir.mkdir(parents=True, exist_ok=True)
        root = work_dir.resolve()
        copied: dict[Path, Path] = {}
        for index, block in enumerate(asset_blocks):
            original = Path(block.image_path)
            try:
                original = original.resolve()
                original.relative_to(root)
            except (OSError, ValueError):
                continue
            if not original.is_file():
                continue
            if original not in copied:
                destination = target_dir / f"{len(copied):03d}_{original.name}"
                shutil.copy2(original, destination)
                copied[original] = destination
            block.image_path = str(copied[original])
        if copied:
            logger.info("已持久化 %d 个解析图片资产: %s", len(copied), target_dir)
    except OSError as exc:
        # 图片资产失败不能让文本解析链路失效；后续按 caption 回退。
        logger.warning("解析图片资产持久化失败: %s", exc)


def _blocks_to_pages(blocks: list[ParsedBlock]) -> list[ParsedPage]:
    """按页聚合 block 文本，保证 pages 视图与 blocks 一致。"""
    by_page: dict[int, list[str]] = {}
    labels: dict[int, str] = {}
    for b in blocks:
        if not b.text:
            continue
        key = b.page or 0
        by_page.setdefault(key, []).append(b.text)
        if b.pdf_page_label and key not in labels:
            labels[key] = b.pdf_page_label
    pages = []
    for k in sorted(by_page):
        pages.append(
            ParsedPage(
                page=None if k == 0 else k,
                text="\n\n".join(by_page[k]),
                pdf_page_label=labels.get(k, ""),
            )
        )
    return pages


def _mineru_json_to_document(content: dict, *, base_dir: Path | None = None) -> ParsedDocument:
    """重组 MinerU content_list 为带 block 结构的 ParsedDocument（结构还原）。

    - 保留 title 层级、table HTML、formula、image 原图路径与说明；
    - 图片/表格说明（caption）就近挂载，无归属的保留为独立 caption block；
    - 页眉页脚保留在 blocks 中（结构还原），正文切片阶段再决定取舍。
    """
    items = content.get("content_list") or []
    raw: list[dict] = []
    for i, block in enumerate(items):
        if isinstance(block, dict):
            item = dict(block)
            item.setdefault("_seq", i)
            raw.append(item)

    # caption 挂载：优先挂到它前面的 image/table（MinerU 输出图在前、说明在后），
    # 无前驱时再找后面的，距离超过 3 个 block 视为无归属。
    targets = [
        (int(b.get("_seq", 0)), b.get("type"))
        for b in raw
        if b.get("type") in ("image", "table")
    ]
    caption_by_target: dict[int, str] = {}
    orphan_captions: list[tuple[int, int | None, str]] = []
    for block in raw:
        if block.get("type") not in _CAPTION_TYPES:
            continue
        text = _caption_text(block)
        if not text:
            continue
        seq = int(block.get("_seq", 0))
        candidates = [t for t in targets if t[0] < seq] or targets
        best_t, best_d = -1, 10**9
        for tseq, _ttype in candidates:
            d = abs(tseq - seq)
            if d < best_d:
                best_t, best_d = tseq, d
        if best_t >= 0 and best_d <= 3:
            caption_by_target[best_t] = text
        else:
            try:
                page = int(block.get("page_idx", 0)) + 1
            except (TypeError, ValueError):
                page = None
            orphan_captions.append((seq, page, text))

    blocks: list[ParsedBlock] = []
    for block in raw:
        try:
            page = int(block.get("page_idx", 0)) + 1
        except (TypeError, ValueError):
            page = None
        btype = str(block.get("type") or "text")
        if btype in {"equation", "equation_interline", "display_formula"}:
            btype = "formula"
        elif btype in {"inline_equation", "inline_formula"}:
            btype = "formula_inline"
        page_label = str(block.get("pdf_page_label") or block.get("page_label") or "").strip()
        seq = int(block.get("_seq", 0))
        order = block.get("order") if block.get("order") is not None else seq

        if btype in _CAPTION_TYPES:
            continue  # 已挂载或进入 orphan_captions
        if btype == "table":
            html = block.get("html") or ""
            text = (block.get("text") or "").strip() or _html_table_to_text(html)
            blocks.append(
                _annotate_block(ParsedBlock(
                    block_type="table",
                    text=text,
                    page=page,
                    html=html,
                    caption=caption_by_target.get(seq, ""),
                    order=order,
                    bbox=_as_bbox(block.get("bbox")),
                    pdf_page_label=page_label,
                ))
            )
        elif btype == "image":
            blocks.append(
                _annotate_block(ParsedBlock(
                    block_type="image",
                    text=(block.get("text") or "").strip(),
                    page=page,
                    image_path=_resolve_image_path(block.get("img_path"), base_dir),
                    caption=caption_by_target.get(seq, ""),
                    order=order,
                    bbox=_as_bbox(block.get("bbox")),
                    pdf_page_label=page_label,
                ))
            )
        elif btype in ("formula", "formula_inline"):
            latex = _formula_latex(block)
            blocks.append(
                _annotate_block(ParsedBlock(
                    block_type=btype,
                    text=f"$${latex}$$" if latex else "",
                    page=page,
                    latex=latex,
                    order=order,
                    bbox=_as_bbox(block.get("bbox")),
                    pdf_page_label=page_label,
                ))
            )
        elif btype == "title":
            try:
                level = int(block.get("level") or 1)
            except (TypeError, ValueError):
                level = 1
            blocks.append(
                _annotate_block(ParsedBlock(
                    block_type="title",
                    text=(block.get("text") or "").strip(),
                    page=page,
                    level=level,
                    order=order,
                    bbox=_as_bbox(block.get("bbox")),
                    pdf_page_label=page_label,
                ))
            )
        else:
            blocks.append(
                _annotate_block(ParsedBlock(
                    block_type=btype,
                    text=_block_text(block).strip(),
                    page=page,
                    order=order,
                    bbox=_as_bbox(block.get("bbox")),
                    pdf_page_label=page_label,
                ))
            )

    for seq, page, text in orphan_captions:
        blocks.append(_annotate_block(ParsedBlock(block_type="caption", text=text, page=page, order=seq)))

    blocks.sort(key=lambda b: (b.page or 0, b.order))
    return ParsedDocument(pages=_blocks_to_pages(blocks), blocks=blocks)


def _mineru_json_to_pages(content: dict) -> ParsedDocument:
    """兼容入口：仅按页重组文本（无结构信息）。"""
    return _mineru_json_to_document(content)


def _mineru_md_to_pages(md_text: str) -> ParsedDocument:
    """无 JSON 时的回退：Markdown 整体作为一页。"""
    text = md_text.strip()
    return ParsedDocument([ParsedPage(page=None, text=text)] if text else [])


def _image_mime(path: Path) -> str:
    return {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".webp": "image/webp",
        ".gif": "image/gif",
        ".bmp": "image/bmp",
    }.get(path.suffix.lower(), "")


def _summarize_image(image_path: str, caption: str, page: int | None) -> str | None:
    """用配置的视觉模型生成图片/图表摘要（多模态视觉理解）。

    未配置 VISUAL_MODEL / 图片缺失 / 调用失败时返回 None，由调用方回退占位，
    不阻断入库主链路。
    """
    p = config.parsing
    if not p.visual_model or not image_path:
        return None
    path = Path(image_path)
    if not path.exists():
        logger.warning("图片不存在，跳过视觉摘要: %s", image_path)
        return None
    mime = _image_mime(path)
    if not mime:
        logger.warning("不支持的图片格式，跳过视觉摘要: %s", path.name)
        return None
    base_url = p.visual_base_url or config.llm.base_url or "https://api.openai.com/v1"
    api_key = p.visual_api_key or config.llm.api_key
    if not api_key:
        logger.warning("VISUAL_MODEL 已配置但无 VISUAL_API_KEY / LLM_API_KEY，跳过视觉摘要")
        return None
    try:
        data_url = (
            f"data:{mime};base64,"
            f"{base64.b64encode(path.read_bytes()).decode('ascii')}"
        )
    except OSError as e:
        logger.warning("读取图片失败，跳过视觉摘要: %s", e)
        return None

    prompt = (
        "请用中文描述这张教学资料图片/图表的核心内容与关键信息，"
        "2-3 句话，适合作为检索摘要，不要输出多余内容。"
    )
    if caption:
        prompt += f"\n图片说明（可能含编号）：{caption}"
    try:
        from openai import OpenAI

        from src.services.http_client import create_openai_http_client

        client = OpenAI(
            api_key=api_key,
            base_url=base_url,
            timeout=p.visual_timeout,
            max_retries=1,
            http_client=create_openai_http_client(p.visual_timeout),
        )
        resp = client.chat.completions.create(
            model=p.visual_model,
            messages=[
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": prompt},
                        {"type": "image_url", "image_url": {"url": data_url}},
                    ],
                }
            ],
            temperature=0.2,
            max_tokens=200,
        )
        text = (resp.choices[0].message.content or "").strip()
        if text:
            logger.info(
                "视觉摘要完成: %s (第%s页)", Path(image_path).name, page or "?"
            )
            return text
    except Exception as e:
        logger.warning("视觉摘要失败，图片按占位文本处理: %s", e)
    return None


def _parse_pdf_mineru(
    path: str,
    *,
    cmd: str = "mineru",
    timeout: int = 0,
    backend: str = "hybrid-engine",
    effort: str = "medium",
    lang: str = "ch",
    formula: bool = True,
    table: bool = True,
    image_analysis: bool = True,
) -> ParsedDocument | None:
    """子进程调用 MinerU CLI，读取输出 Markdown/JSON 重组为 ParsedDocument。

    timeout<=0 表示不限时；任何失败返回 None，由调用方回退现有链路。
    """
    if not _mineru_available(cmd):
        return None
    tmp = Path(tempfile.mkdtemp(prefix="mineru_"))
    try:
        args = [
            cmd,
            "-p", str(path),
            "-o", str(tmp),
            "-m", "auto",
            "-b", backend,
            "-l", lang,
            "-f", str(formula).lower(),
            "-t", str(table).lower(),
            "--image-analysis", str(image_analysis).lower(),
        ]
        # --effort 仅 hybrid backend 支持；其他 backend 保持兼容。
        if backend in ("hybrid-engine", "hybrid-http-client"):
            args.extend(["--effort", effort])
        logger.info(
            "MinerU 解析开始: %s (backend=%s effort=%s timeout=%s)",
            Path(path).name,
            backend,
            effort,
            timeout or "无",
        )
        proc, stdout, stderr = _run_capture(
            args,
            timeout=None if timeout <= 0 else timeout,
        )
        if proc.returncode != 0:
            logger.warning(
                "MinerU 退出码 %s: %s",
                proc.returncode,
                _command_error_tail(stdout, stderr),
            )
            return None

        md_files = sorted(tmp.rglob("*.md"))
        json_files = sorted(tmp.rglob("*.json"))
        md_doc = None
        if md_files:
            md_doc = _mineru_md_to_pages(
                md_files[0].read_text(encoding="utf-8", errors="replace")
            )
        for jf in json_files:
            try:
                content = json.loads(jf.read_text(encoding="utf-8", errors="replace"))
            except Exception:
                continue
            if isinstance(content, dict) and isinstance(content.get("content_list"), list):
                doc = _mineru_json_to_document(content, base_dir=jf.parent)
                if doc.pages:
                    _persist_parser_assets(doc, work_dir=tmp, source_path=path)
                    logger.info(
                        "MinerU 解析完成: %s, %d 页 / %d 块 (json)",
                        Path(path).name,
                        len(doc.pages),
                        len(doc.blocks),
                    )
                    return doc
        if md_doc and md_doc.pages:
            logger.info("MinerU 解析完成: %s, 回退 Markdown", Path(path).name)
            return md_doc
        logger.warning("MinerU 未产出可用文本: %s", Path(path).name)
        return None
    except subprocess.TimeoutExpired:
        logger.warning("MinerU 解析超时（%s 秒）: %s", timeout, Path(path).name)
        return None
    except Exception as e:
        logger.warning("MinerU 解析异常: %s", e)
        return None
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _markpdfdown_available(cmd: str) -> bool:
    """外部增强解析器命令是否可调用。"""
    return bool(cmd and shutil.which(cmd))


def _markpdfdown_command(
    cmd: str,
    args_template: str,
    *,
    path: str,
    output_dir: Path,
    output_file: Path,
) -> list[str] | None:
    """按显式模板构建外部解析器命令，禁止 shell 执行与猜测 CLI 参数。

    {output} 指临时目录，适用于目录型 CLI；{output_file} 指临时 .md 文件，
    适用于 MarkPDFdown 的 ``--output`` 文件参数。
    """
    if "{input}" not in args_template or (
        "{output}" not in args_template and "{output_file}" not in args_template
    ):
        logger.warning(
            "MARKPDFDOWN_ARGS 必须包含 {input} 与 {output}/{output_file} 占位符，跳过增强解析器"
        )
        return None
    try:
        # 命令最终以 list 传给 subprocess，不需要保留引号；posix=True 可统一
        # 去除模板中的引号（Windows 下 shlex 的非 POSIX 模式会把引号当作字符）。
        tokens = shlex.split(args_template, posix=True)
    except ValueError as exc:
        logger.warning("MARKPDFDOWN_ARGS 格式错误，跳过增强解析器: %s", exc)
        return None
    if not tokens:
        logger.warning("MARKPDFDOWN_ARGS 为空，跳过增强解析器")
        return None
    replacements = {
        "{input}": str(Path(path).resolve()),
        "{output}": str(output_dir),
        "{output_file}": str(output_file),
    }
    return [
        cmd,
        *[
            token.replace("{input}", replacements["{input}"])
            .replace("{output_file}", replacements["{output_file}"])
            .replace("{output}", replacements["{output}"])
            for token in tokens
        ],
    ]


def _generic_json_to_document(content: object, *, base_dir: Path) -> ParsedDocument | None:
    """兼容外部解析器常见 JSON 输出；优先复用 MinerU 的 content_list 结构。"""
    if isinstance(content, dict) and isinstance(content.get("content_list"), list):
        return _mineru_json_to_document(content, base_dir=base_dir)
    if isinstance(content, dict):
        raw_pages = content.get("pages") or (content.get("data") or {}).get("pages")
    else:
        raw_pages = None
    if not isinstance(raw_pages, list):
        return None
    pages: list[ParsedPage] = []
    blocks: list[ParsedBlock] = []
    for index, item in enumerate(raw_pages, 1):
        if not isinstance(item, dict):
            continue
        page = _page_num(item) or index
        text = str(item.get("text") or item.get("content") or item.get("markdown") or "").strip()
        if text:
            pages.append(ParsedPage(page=page, text=text))
        for order, block in enumerate(item.get("blocks") or []):
            if not isinstance(block, dict):
                continue
            block_type = str(block.get("type") or "text")
            if block_type in {"equation", "equation_interline", "display_formula"}:
                block_type = "formula"
            elif block_type in {"inline_equation", "inline_formula"}:
                block_type = "formula_inline"
            if block_type in {"formula", "formula_inline"}:
                latex = _formula_latex(block)
                block_text = f"$${latex}$$" if latex else ""
            else:
                block_text = str(block.get("text") or block.get("content") or "").strip()
            if block_text:
                blocks.append(
                    _annotate_block(ParsedBlock(
                        block_type=block_type,
                        text=block_text,
                        page=page,
                        order=order,
                        pdf_page_label=str(item.get("page_label") or "").strip(),
                    ))
                )
    if blocks and not pages:
        pages = _blocks_to_pages(blocks)
    return ParsedDocument(pages=pages, blocks=blocks) if pages else None


def _parse_pdf_markpdfdown(
    path: str, *, cmd: str, args_template: str, timeout: int
) -> ParsedDocument | None:
    """运行可配置的 MarkPDFdown 适配器并读取 Markdown/JSON 输出。

    项目不假设第三方 CLI 的参数名；使用者必须在 MARKPDFDOWN_ARGS 指定
    带 {input}/{output} 的准确命令模板。任何异常均回退内置解析链。
    """
    if not _markpdfdown_available(cmd):
        logger.warning("MarkPDFdown 命令不可用: %s", cmd or "<未配置>")
        return None
    tmp = Path(tempfile.mkdtemp(prefix="markpdfdown_"))
    try:
        args = _markpdfdown_command(
            cmd,
            args_template,
            path=path,
            output_dir=tmp,
            output_file=tmp / "document.md",
        )
        if not args:
            return None
        logger.info("MarkPDFdown 增强解析开始: %s", Path(path).name)
        proc, stdout, stderr = _run_capture(
            args, timeout=None if timeout <= 0 else timeout
        )
        if proc.returncode != 0:
            logger.warning(
                "MarkPDFdown 退出码 %s: %s",
                proc.returncode,
                _command_error_tail(stdout, stderr),
            )
            return None
        for json_path in sorted(tmp.rglob("*.json")):
            try:
                content = json.loads(json_path.read_text(encoding="utf-8", errors="replace"))
            except (OSError, json.JSONDecodeError):
                continue
            if doc := _generic_json_to_document(content, base_dir=json_path.parent):
                _persist_parser_assets(doc, work_dir=tmp, source_path=path)
                logger.info("MarkPDFdown 解析完成: %s (json)", Path(path).name)
                return doc
        markdown_files = [*tmp.rglob("*.md"), *tmp.rglob("*.markdown")]
        if markdown_files:
            best = max(markdown_files, key=lambda item: item.stat().st_size)
            doc = _mineru_md_to_pages(best.read_text(encoding="utf-8", errors="replace"))
            if doc.pages:
                logger.info("MarkPDFdown 解析完成: %s (markdown)", Path(path).name)
                return doc
        logger.warning("MarkPDFdown 未产出可用 JSON/Markdown: %s", Path(path).name)
        return None
    except subprocess.TimeoutExpired:
        logger.warning("MarkPDFdown 解析超时（%s 秒）: %s", timeout, Path(path).name)
        return None
    except Exception as exc:
        logger.warning("MarkPDFdown 解析异常: %s", exc)
        return None
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _parse_pdf(path: str) -> ParsedDocument:
    p = config.parsing
    pdf_kind = _pdf_kind(path)
    expected_pages = _pdf_page_count(path)
    logger.info("PDF 类型判断: %s -> %s", Path(path).name, pdf_kind)

    best_doc: ParsedDocument | None = None
    best_quality = _assess_pdf_quality(None, expected_pages=expected_pages)

    def consider(doc: ParsedDocument | None, source: str) -> bool:
        nonlocal best_doc, best_quality
        quality = _assess_pdf_quality(doc, expected_pages=expected_pages)
        _log_pdf_quality(source, quality)
        if doc is not None:
            doc.parser_name = source
            doc.parse_quality = quality.score
        if doc and doc.full_text.strip() and quality.score > best_quality.score:
            best_doc, best_quality = doc, quality
        return bool(doc and doc.full_text.strip() and quality.score >= p.pdf_quality_threshold)

    # 显式指定或 auto + 开关时，先调用可配置的外部增强解析器。
    want_markpdfdown = p.pdf_parser == "markpdfdown" or (
        p.pdf_parser == "auto" and p.markpdfdown_enabled
    )
    if want_markpdfdown:
        if consider(
            _parse_pdf_markpdfdown(
                path,
                cmd=p.markpdfdown_cmd,
                args_template=p.markpdfdown_args,
                timeout=p.markpdfdown_timeout,
            ),
            "markpdfdown",
        ):
            return best_doc  # type: ignore[return-value]

    # MarkPDFdown 是增强候选而不是单点依赖；显式选它失败时仍回退 MinerU。
    if p.pdf_parser in ("mineru", "auto", "markpdfdown"):
        want_mineru = p.pdf_parser in ("mineru", "markpdfdown") or pdf_kind in ("scanned", "mixed")
        if want_mineru:
            if _mineru_available(p.mineru_cmd):
                doc = _parse_pdf_mineru(
                    path,
                    cmd=p.mineru_cmd,
                    timeout=p.mineru_timeout,
                    backend=p.mineru_backend,
                    effort=p.mineru_effort,
                    lang=p.mineru_lang,
                    formula=p.mineru_formula,
                    table=p.mineru_table,
                    image_analysis=p.mineru_image_analysis,
                )
                if consider(doc, f"mineru:{p.mineru_effort}"):
                    return best_doc  # type: ignore[return-value]
                if p.mineru_retry_high and p.mineru_effort != "high":
                    high_doc = _parse_pdf_mineru(
                        path,
                        cmd=p.mineru_cmd,
                        timeout=p.mineru_timeout,
                        backend=p.mineru_backend,
                        effort="high",
                        lang=p.mineru_lang,
                        formula=p.mineru_formula,
                        table=p.mineru_table,
                        image_analysis=True,
                    )
                    if consider(high_doc, "mineru:high"):
                        return best_doc  # type: ignore[return-value]
            elif p.pdf_parser == "mineru":
                logger.warning(
                    "PDF_PARSER=mineru 但找不到命令 %s，回退现有链路", p.mineru_cmd
                )

    # 原生文本 PDF：直接提取
    try:
        doc = _parse_pdf_pymupdf4llm(path)
        doc = _enrich_document_with_formulas(doc, path)
        if consider(doc, "pymupdf4llm"):
            return best_doc  # type: ignore[return-value]
        if p.pdf_use_ocr and not p.pdf_force_ocr:
            logger.info("PDF 空文本，OCR 重试: %s", Path(path).name)
            doc = _parse_pdf_pymupdf4llm(path, force_ocr=True)
            doc = _enrich_document_with_formulas(doc, path)
            if consider(doc, "pymupdf4llm:ocr"):
                return best_doc  # type: ignore[return-value]
    except Exception as e:
        logger.warning("pymupdf4llm 失败，回退 fitz: %s", e)

    doc = _parse_pdf_fitz(path)
    if consider(doc, "fitz"):
        return best_doc  # type: ignore[return-value]

    if p.pdf_use_ocr and not p.pdf_force_ocr:
        try:
            ocr = _parse_pdf_pymupdf4llm(path, force_ocr=True)
            ocr = _enrich_document_with_formulas(ocr, path)
            if consider(ocr, "pymupdf4llm:ocr-final"):
                return best_doc  # type: ignore[return-value]
        except Exception as e:
            logger.warning("PDF OCR 失败: %s", e)
    if best_doc:
        logger.warning(
            "PDF 所有候选均未达到质量阈值 %.2f，保留最佳结果（%.2f）",
            p.pdf_quality_threshold,
            best_quality.score,
        )
        return best_doc
    return doc


def _parse_txt(path: str) -> str:
    data = Path(path).read_bytes()
    for enc in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    # ponytail: 怪编码替换非法字节，不阻断
    return data.decode("utf-8", errors="replace")


def _parse_plain(path: str) -> ParsedDocument:
    text = _parse_txt(path).strip()
    return ParsedDocument([ParsedPage(page=None, text=text)] if text else [])


def _table_to_text(table) -> str:
    rows: list[str] = []
    for row in table.rows:
        seen: set[int] = set()
        cells: list[str] = []
        for cell in row.cells:
            key = id(cell._tc)
            if key in seen:
                continue
            seen.add(key)
            if t := cell.text.strip():
                cells.append(t)
            for nested in cell.tables:
                if nt := _table_to_text(nested):
                    cells.append(nt)
        if cells:
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _iter_block_items(parent):
    from docx.document import Document as DocxDocument
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = parent.element.body if isinstance(parent, DocxDocument) else parent._element
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def _story_parts(container) -> list[str]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parts: list[str] = []
    for block in _iter_block_items(container):
        if isinstance(block, Paragraph):
            if t := (block.text or "").strip():
                parts.append(t)
        elif isinstance(block, Table):
            if t := _table_to_text(block):
                parts.append(t)
    return parts


def _header_footer_parts(doc) -> list[str]:
    # 单节异常跳过，避免整份 docx 失败
    seen: set[str] = set()
    out: list[str] = []
    for section in doc.sections:
        for part in (section.header, section.footer):
            try:
                block = "\n".join(_story_parts(part)).strip()
            except Exception:
                continue
            if block and block not in seen:
                seen.add(block)
                out.append(block)
    return out


def _textbox_parts(doc) -> list[str]:
    from docx.oxml.ns import qn

    parts: list[str] = []
    for txbx in doc.element.body.iter(qn("w:txbxContent")):
        texts = [(n.text or "").strip() for n in txbx.iter(qn("w:t")) if (n.text or "").strip()]
        if texts:
            parts.append("\n".join(texts))
    return parts


def _parse_docx(path: str) -> ParsedDocument:
    from docx import Document

    doc = Document(path)
    parts = _header_footer_parts(doc) + _story_parts(doc) + _textbox_parts(doc)
    text = "\n\n".join(parts)
    return ParsedDocument([ParsedPage(page=None, text=text)] if text else [])


def _find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        if found := shutil.which(name):
            return found
    if os.name == "nt":
        for exe in ("soffice.com", "soffice.exe"):
            for pattern in (
                rf"C:\Program Files\LibreOffice\program\{exe}",
                rf"C:\Program Files (x86)\LibreOffice\program\{exe}",
            ):
                if Path(pattern).is_file():
                    return pattern
    return None


def _convert_doc_to_docx_soffice(path: str) -> Path | None:
    soffice = _find_soffice()
    if not soffice:
        return None

    out_dir = Path(tempfile.mkdtemp(prefix="exam_doc_"))
    dest: Path | None = None
    try:
        proc, _stdout, stderr = _run_capture(
            [soffice, "--headless", "--convert-to", "docx", "--outdir", str(out_dir), path],
            timeout=120,
        )
        if proc.returncode != 0:
            logger.warning(
                "LibreOffice 退出码 %s: %s", proc.returncode, stderr.strip()[:300]
            )
            return None
        converted = out_dir / f"{Path(path).stem}.docx"
        if not converted.is_file():
            return None
        fd, dest_name = tempfile.mkstemp(suffix=".docx", prefix="exam_doc_")
        os.close(fd)
        dest = Path(dest_name)
        shutil.copy2(converted, dest)
        return dest
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError) as e:
        logger.warning("LibreOffice 转换失败: %s", e)
        if dest is not None:
            dest.unlink(missing_ok=True)
        return None
    finally:
        shutil.rmtree(out_dir, ignore_errors=True)


def _convert_doc_to_docx_word(path: str) -> Path | None:
    """Windows：无 LibreOffice 时用本机 Word COM 另存为 .docx。"""
    if os.name != "nt":
        return None
    src = str(Path(path).resolve())
    fd, dest_name = tempfile.mkstemp(suffix=".docx", prefix="exam_doc_")
    os.close(fd)
    dest = Path(dest_name)
    dest.unlink(missing_ok=True)
    src_ps = src.replace("'", "''")
    dest_ps = str(dest.resolve()).replace("'", "''")
    ps = (
        "$ErrorActionPreference='Stop'; "
        "$word=New-Object -ComObject Word.Application; "
        "$word.Visible=$false; "
        f"$doc=$word.Documents.Open('{src_ps}'); "
        f"$doc.SaveAs([ref]'{dest_ps}',[ref]16); "
        "$doc.Close(); $word.Quit(); "
        "[System.Runtime.InteropServices.Marshal]::ReleaseComObject($word)|Out-Null"
    )
    try:
        proc, _stdout, stderr = _run_capture(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", ps],
            timeout=120,
        )
        if proc.returncode != 0:
            logger.warning(
                "Word COM 退出码 %s: %s", proc.returncode, stderr.strip()[:300]
            )
            return None
        if dest.is_file() and dest.stat().st_size > 0:
            return dest
        dest.unlink(missing_ok=True)
        return None
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError) as e:
        logger.warning("Word COM 转换失败: %s", e)
        dest.unlink(missing_ok=True)
        return None


def _convert_doc_to_docx(path: str) -> Path | None:
    return _convert_doc_to_docx_soffice(path) or _convert_doc_to_docx_word(path)


def _parse_doc(path: str) -> ParsedDocument:
    docx_tmp = _convert_doc_to_docx(path)
    if docx_tmp:
        try:
            return _parse_docx(str(docx_tmp))
        finally:
            docx_tmp.unlink(missing_ok=True)
    raise BadRequestException(
        "无法解析 .doc 文件。请安装 LibreOffice 或 Microsoft Word，或将文件另存为 .docx 后重试。"
    )


def _shape_texts(shape) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
        return [t for child in shape.shapes for t in _shape_texts(child)]

    texts: list[str] = []
    if getattr(shape, "has_table", False):
        rows = []
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            texts.append("\n".join(rows))
    if getattr(shape, "has_text_frame", False):
        if t := shape.text_frame.text.strip():
            texts.append(t)
    elif hasattr(shape, "text") and (t := (shape.text or "").strip()):
        texts.append(t)
    return texts


def _convert_pptx_to_pdf_soffice(path: str) -> Path | None:
    """LibreOffice headless 将 PPTX 转为 PDF，可覆盖 SmartArt/图表/母版文本。"""
    soffice = _find_soffice()
    if not soffice:
        return None

    out_dir = Path(tempfile.mkdtemp(prefix="exam_pptx_"))
    profile_dir = Path(tempfile.mkdtemp(prefix="exam_lo_profile_"))
    dest: Path | None = None
    try:
        user_install = "file:///" + str(profile_dir).replace("\\", "/")
        proc, _stdout, stderr = _run_capture(
            [
                soffice,
                "--headless",
                "--norestore",
                "--nolockcheck",
                "--convert-to",
                "pdf",
                "--outdir",
                str(out_dir),
                f"-env:UserInstallation={user_install}",
                path,
            ],
            timeout=180,
        )
        if proc.returncode != 0:
            logger.warning(
                "LibreOffice 退出码 %s: %s", proc.returncode, stderr.strip()[:300]
            )
            return None
        converted = out_dir / f"{Path(path).stem}.pdf"
        if not converted.is_file():
            return None
        fd, dest_name = tempfile.mkstemp(suffix=".pdf", prefix="exam_pptx_")
        os.close(fd)
        dest = Path(dest_name)
        shutil.copy2(converted, dest)
        return dest
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError) as e:
        logger.warning("LibreOffice PPTX 转 PDF 失败: %s", e)
        if dest is not None:
            dest.unlink(missing_ok=True)
        return None
    finally:
        shutil.rmtree(out_dir, ignore_errors=True)
        shutil.rmtree(profile_dir, ignore_errors=True)


def _convert_pptx_to_pdf_powerpoint(path: str) -> Path | None:
    """Windows：无 LibreOffice 时用本机 PowerPoint COM 导出 PDF（ppSaveAsPDF=32）。"""
    if os.name != "nt":
        return None
    src = str(Path(path).resolve())
    fd, dest_name = tempfile.mkstemp(suffix=".pdf", prefix="exam_pptx_")
    os.close(fd)
    dest = Path(dest_name)
    dest.unlink(missing_ok=True)
    src_ps = src.replace("'", "''")
    dest_ps = str(dest.resolve()).replace("'", "''")
    ps = (
        "$ErrorActionPreference='Stop'; "
        "$pp=New-Object -ComObject PowerPoint.Application; "
        "try { "
        f"$pres=$pp.Presentations.Open('{src_ps}',-1,0,0); "
        f"$pres.SaveAs([ref]'{dest_ps}',[ref]32); "
        "$pres.Close(); "
        "} finally { "
        "$pp.Quit(); "
        "[System.Runtime.InteropServices.Marshal]::ReleaseComObject($pp)|Out-Null "
        "}"
    )
    try:
        subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", ps],
            check=True,
            capture_output=True,
            text=True,
            timeout=180,
        )
        if dest.is_file() and dest.stat().st_size > 0:
            return dest
        dest.unlink(missing_ok=True)
        return None
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError) as e:
        logger.warning("PowerPoint COM 导出 PDF 失败: %s", e)
        dest.unlink(missing_ok=True)
        return None


def _convert_pptx_to_pdf(path: str) -> Path | None:
    return _convert_pptx_to_pdf_soffice(path) or _convert_pptx_to_pdf_powerpoint(path)


def _parse_pptx_direct(path: str) -> ParsedDocument:
    from pptx import Presentation

    pages: list[ParsedPage] = []
    for i, slide in enumerate(Presentation(path).slides, 1):
        texts = [t for shape in slide.shapes for t in _shape_texts(shape)]
        if slide.has_notes_slide:
            if notes := slide.notes_slide.notes_text_frame.text.strip():
                texts.append(f"[备注]\n{notes}")
        if texts:
            pages.append(ParsedPage(page=i, text="\n".join(texts)))
    return ParsedDocument(pages)


def _tesseract_available() -> bool:
    if shutil.which("tesseract"):
        return True
    if os.name == "nt" and Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe").is_file():
        return True
    return False


def _clean_pdf_markdown(text: str) -> str:
    """清理 pymupdf4llm 输出的 markdown 杂质（图片占位符、标题标记等）。"""
    text = re.sub(r"\*\*==>.*?<==\*\*", "", text, flags=re.S)
    text = re.sub(r"^\s*#+\s*$", "", text, flags=re.M)
    text = re.sub(r"^\s*#{1,6}\s+", "", text, flags=re.M)
    text = text.replace("**", "").replace("__", "")
    text = re.sub(r"^\s*\|?[\s:|-]+\|?\s*$", "", text, flags=re.M)
    lines = [ln.rstrip() for ln in text.splitlines() if ln.strip()]
    return "\n".join(lines).strip()


def _merge_pages_by_length(a: list[ParsedPage], b: list[ParsedPage]) -> ParsedDocument:
    """同一 PPT 的两种提取结果按页择优合并（PDF 页号与 PPT 页号一致）。"""
    pages: list[ParsedPage] = []
    for i in range(max(len(a), len(b))):
        pa = a[i] if i < len(a) else None
        pb = b[i] if i < len(b) else None
        if pa and pb:
            pages.append(pa if len(pa.text) >= len(pb.text) else pb)
        else:
            pages.append(pa or pb)
    return ParsedDocument([p for p in pages if p.text and p.text.strip()])


def _parse_pptx(path: str) -> ParsedDocument:
    """优先转 PDF 走 pymupdf4llm 管线，清理杂质后与 python-pptx 按页择优合并；
    图片型 PPT 自动尝试 OCR（需系统安装 Tesseract）。"""
    direct = _parse_pptx_direct(path)
    pdf_tmp = _convert_pptx_to_pdf(path)
    if not pdf_tmp:
        return direct
    try:
        doc = _parse_pdf(str(pdf_tmp))
        pdf_pages = [
            ParsedPage(page=p.page, text=_clean_pdf_markdown(p.text))
            for p in doc.pages
            if _clean_pdf_markdown(p.text).strip()
        ]
        direct_chars = sum(len(p.text) for p in direct.pages)
        pdf_chars = sum(len(p.text) for p in pdf_pages)
        if pdf_chars < max(direct_chars, 1) * 0.5 and _tesseract_available():
            try:
                ocr_doc = _parse_pdf_pymupdf4llm(str(pdf_tmp), force_ocr=True)
                ocr_pages = [
                    ParsedPage(page=p.page, text=_clean_pdf_markdown(p.text))
                    for p in ocr_doc.pages
                    if _clean_pdf_markdown(p.text).strip()
                ]
                ocr_chars = sum(len(p.text) for p in ocr_pages)
                if ocr_chars > pdf_chars:
                    logger.info("PPTX 图片 OCR 生效: %s -> %d 字", Path(path).name, ocr_chars)
                    pdf_pages = ocr_pages
            except Exception as e:
                logger.warning("PPTX 图片 OCR 失败: %s", e)
        return _merge_pages_by_length(direct.pages, pdf_pages)
    except Exception as e:
        logger.warning("PPTX 转 PDF 解析失败，回退 python-pptx: %s", e)
        return direct
    finally:
        pdf_tmp.unlink(missing_ok=True)


_PARSERS = {
    ".pdf": _parse_pdf,
    ".txt": _parse_plain,
    ".md": _parse_plain,
    ".doc": _parse_doc,
    ".docx": _parse_docx,
    ".pptx": _parse_pptx,
}


def parse_file(path: str) -> ParsedDocument:
    ext = Path(path).suffix.lower()
    if ext not in _PARSERS:
        raise UnsupportedFormatException(
            f"不支持的文件格式: {ext}，仅接受 PDF/TXT/MD/DOC/DOCX/PPTX"
        )
    if not os.path.exists(path):
        raise BadRequestException(f"文件不存在: {path}")

    try:
        doc = _PARSERS[ext](path)
    except (UnsupportedFormatException, BadRequestException):
        raise
    except Exception as e:
        logger.error("文件解析失败 (%s): %s", path, e)
        raise BadRequestException(f"文件解析失败: {e}") from e

    if not doc.full_text.strip():
        raise BadRequestException("文件内容为空，无法入库")

    if not doc.parser_name:
        doc.parser_name = ext.lstrip(".")

    logger.info(
        "解析完成: %s parser=%s quality=%.2f pages=%d chars=%d",
        Path(path).name,
        doc.parser_name,
        doc.parse_quality,
        len(doc.pages),
        len(doc.full_text),
    )
    return doc
